"""Export Excel / JSON des résultats."""
from __future__ import annotations

import io
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERDICT_FILL = {
    "CONFORME": "C6EFCE", "CONFORME_TOLERE": "E2F0D9", "A_VERIFIER": "FFEB9C",
    "NON_CONFORME": "F8CBAD", "PIECE_MANQUANTE": "D9D9D9", "PIECE_INUTILISABLE": "D9D9D9",
}

HEAD = ["Ligne CSV", "Contrat", "Rôle", "Verdict",
        "Nom saisi", "Prénoms saisis", "Nom complet saisi", "Numéro saisi",
        "Nom lu", "Prénoms lus", "Numéro lu", "Type lu",
        "Statut nom", "Score nom", "Statut numéro", "Score numéro",
        "Motifs", "Signaux", "Scan(s)", "Sources (identité / numéro)"]


def flatten(r: dict) -> list:
    c, e, k = r.get("csv", {}), r.get("extracted", {}), r.get("checks", {})
    n, p = k.get("nom_complet", {}), k.get("numero_piece", {})
    return [
        r["row_index"] + 2, r["code"], r["role"], r["verdict"],
        c.get("nom"), c.get("prenoms"), c.get("nom_complet"), c.get("numero"),
        e.get("nom"), e.get("prenoms"), e.get("numero_piece"), e.get("type_piece"),
        n.get("status"), n.get("score"), p.get("status"), p.get("score"),
        " ; ".join(r.get("reasons", [])), " ; ".join(r.get("flags", [])),
        " ; ".join(s["fichier"] for s in r.get("scans", [])),
        f"{e.get('source_identite') or ''} / {e.get('source_numero') or ''}".strip(" /"),
    ]


def to_xlsx(results: list[dict], orphans: list[dict], summary: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats"
    ws.append(HEAD)
    for r in results:
        ws.append(flatten(r))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        fill = VERDICT_FILL.get(row[3].value)
        if fill:
            row[3].fill = PatternFill("solid", fgColor=fill)
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [9, 26, 14, 20, 20, 20, 28, 18, 20, 20, 18, 14, 12, 9, 12, 9, 50, 36, 46, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    wo = wb.create_sheet("Scans non rattachés")
    wo.append(["Fichier", "Contrat (nom de fichier)", "Indice nom (fichier)", "Type lu", "Nom lu", "Prénoms lus",
               "Numéro lu", "Face", "Ligne CSV la plus proche", "Score"])
    for o in orphans:
        ex = o.get("extraction") or {}
        b = o.get("plus_proche") or {}
        wo.append([o["fichier"], o["code"], o["indice_nom"], ex.get("type_piece"), ex.get("nom"), ex.get("prenoms"),
                   ex.get("numero_piece"), ex.get("face"), b.get("nom"), b.get("score")])
    for cell in wo[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate([48, 26, 30, 14, 22, 22, 18, 10, 30, 8], 1):
        wo.column_dimensions[get_column_letter(i)].width = w

    wd = wb.create_sheet("Synthèse")
    wd.append(["Verdict", "Nombre"])
    for k, v in summary.items():
        wd.append([k, v])
    wd.column_dimensions["A"].width = 24

    wj = wb.create_sheet("Extractions (JSON)")
    wj.append(["Ligne CSV", "Fichier", "JSON extrait"])
    for r in results:
        for s in r.get("scans", []):
            wj.append([r["row_index"] + 2, s["fichier"], json.dumps(s.get("extraction"), ensure_ascii=False)])
    wj.column_dimensions["B"].width = 50
    wj.column_dimensions["C"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
