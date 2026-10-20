"""Exports d'une campagne : Excel des remarques (13 colonnes), classeurs RAS/REAMRQUE/Global, dossier complet avec pièces."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .contracts import COLUMNS, evaluate, export_row, norm
from .identity_check import ROLE_CODES, basename, resolve_scan

_CODE = re.compile(r"CT-\d{3}-\d{3}-\d{6}-\d{2}-[a-z]+", re.I)
HEAD_FILL = PatternFill("solid", fgColor="163F38")


def _cell(ws, row: int, col: int, value):
    c = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] in "=+-@":
        c.data_type = "s"          # jamais interprété comme formule
    return c


def make_workbook(contracts: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contrats"
    ws.append(COLUMNS)
    for c in contracts:
        for j, v in enumerate(export_row(c), 1):
            _cell(ws, ws.max_row + 1 if j == 1 else ws.max_row, j, v)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = HEAD_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(64 + len(COLUMNS))}{max(1, ws.max_row)}"
    for i, col in enumerate(ws.columns):
        ws.column_dimensions[chr(65 + i)].width = 48 if i > 5 else 25
        for cell in col:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def output_names(contracts: list[dict]) -> dict[str, str]:
    used: set[str] = set()
    names: dict[str, str] = {}
    for source in dict.fromkeys(c["source"] for c in contracts):
        base = basename(source)
        base = re.sub(r"\.(csv|xlsx?|json)$", "", base, flags=re.I)
        base = re.sub(r"ALL_TABLES_EXPORT", "", base, flags=re.I)
        base = re.sub(r"^[_ .-]+|[_ .-]+$", "", base)
        base = re.sub(r"[<>:\"/\\|?*\x00-\x1f]", "_", base).rstrip(". ") or "Fichier"
        name, n = f"Rque_contrat_{base}.xlsx", 2
        while name.lower() in used:
            name = f"Rque_contrat_{base}_{n}.xlsx"
            n += 1
        used.add(name.lower())
        names[source] = name
    return names


def export_remarks(contracts: list[dict]) -> tuple[str, bytes, str, int]:
    """(nom de fichier, contenu, type MIME, nombre de contrats avec remarques)."""
    rows = [c for c in map(evaluate, contracts) if c["observations"]]
    names = output_names(contracts)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if len(names) <= 1:
        return (next(iter(names.values()), "Rque_contrat_Global.xlsx"), make_workbook(rows), xlsx, len(rows))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for source, name in names.items():
            z.writestr(name, make_workbook([c for c in rows if c["source"] == source]))
        z.writestr("Global/Rque_contrat_Global.xlsx", make_workbook(rows))
    return "Rque_contrat_Resultats.zip", buf.getvalue(), "application/zip", len(rows)


# ------------------------------------------------------------------ pièces
def _safe(v: str) -> str:
    return re.sub(r"[^a-zA-Z0-9_-]+", "_", v)[:160] or "NON_RENSEIGNE"


def plan_pieces(contracts: list[dict], names: list[str], aliases: dict[str, str] | None = None) -> list[dict]:
    entries, used, counts = [], set(), {}
    seen: set[tuple[str, str]] = set()
    aliases = aliases or {}

    def add(contract, original, actors):
        if (contract, original) in seen:
            return
        seen.add((contract, original))
        distinct = list({f"{a['role']}::{a['name']}": a for a in actors}.values())
        actor = distinct[0] if len(distinct) == 1 else None
        role = actor["role"] if actor else "Qualité à confirmer"
        abbr = ROLE_CODES.get(role, "QUAL_A_CONFIRMER")
        person = actor["name"] if actor else "Personne à confirmer"
        m = re.search(r"\.(jpg|jpeg|png|webp|pdf)$", original, re.I)
        ext = m.group(0).lower() if m else ".jpg"
        base = f"{_safe(contract)}_{abbr}_{_safe(norm(person))}"
        counts[base] = counts.get(base, 0) + 1
        entries.append({"contract": contract, "role": role, "abbr": abbr, "person": person, "original": original,
                        "path": f"PIECE SCANNEE/{_safe(contract)}/{base}_{counts[base]:02d}{ext}"})
        used.add(original)

    def own(a):
        return list(dict.fromkeys([*(resolve_scan(p, names, aliases) or p for p in a.get("pieces", [])),
                                   *(a.get("automatic") or {}).get("pieces", [])]))

    for c in contracts:
        refs = list(dict.fromkeys(p for a in c["actors"] for p in own(a)))
        for original in refs:
            add(c["number"], original, [a for a in c["actors"] if original in own(a)])
    by_number = {c["number"].upper(): c for c in contracts}
    for original in [n for n in names if n not in used]:
        m = _CODE.search(original)
        code = m.group(0).upper() if m else None
        c = by_number.get(code) if code else None
        add(c["number"] if c else (code or "CONTRAT_A_CONFIRMER"), original, [])
    return entries


def _csv_cell(s: str) -> str:
    s = "" if s is None else str(s)
    return '"' + ("'" if s[:1] in "=+@-" else "") + s.replace('"', '""') + '"'


def ocr_workbook(contracts: list[dict], ocr: dict) -> bytes:
    heads = ["Nom_fichier", "Page", "Type_document", "Nom_scan", "Prenoms_scan", "Numero_piece_scan", "Date_naissance_scan",
             "Date_expiration_scan", "Sexe_scan", "Numero_contrat", "Qualite", "Nom_complet_csv", "Numero_piece_csv",
             "Date_naissance_csv", "Score_nom_complet", "Score_numero", "Score_date_naissance", "Score_global", "Statut",
             "Observation", "Date_traitement", "Confiance_OCR", "Tentatives", "Texte_OCR", "Valeurs_normalisées"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Pièces et pages"
    ws.append(heads)
    owner: dict[str, tuple[dict, dict]] = {}
    for c in contracts:
        for a in c["actors"]:
            for p in (a.get("automatic") or {}).get("pieces", []):
                owner.setdefault(p, (c, a))
    from .normalize import name_tokens, norm_number
    for key in sorted(ocr):
        rec = ocr[key]
        ex = rec.get("extraction") or {}
        fname, _, page = key.partition("#p")
        c, a = owner.get(fname, (None, None))
        chk = (a or {}).get("automatic") or {}
        status = {"validé": "CONFORME", "non conforme": "NON CONFORME"}.get(chk.get("status"), "À VÉRIFIER")
        obs = "RAS" if status == "CONFORME" else (chk.get("detail") or ex.get("erreur") or
                                                 "Aucune correspondance établie dans le CSV" if not a else "")
        confs = [v for v in (ex.get("confiance") or {}).values() if v is not None]
        row = [fname, int(page) if page else 1, ex.get("type_piece"), ex.get("nom"), ex.get("prenoms"),
               ex.get("numero_piece"), ex.get("date_naissance"), None, None, c["number"] if c else "",
               a["role"] if a else "", a["name"] if a else "", a["id"] if a else "", (a or {}).get("birthDate", ""),
               chk.get("score") if a else None, None, None, chk.get("score") if a else None, status, obs,
               rec.get("at", ""), round(100 * max(confs)) if confs else 0, 1, ex.get("remarques") or "",
               json.dumps({"nom": " ".join(name_tokens(" ".join(x for x in (ex.get("nom"), ex.get("prenoms")) if x))),
                           "numero": norm_number(ex.get("numero_piece") or "")}, ensure_ascii=False)]
        for j, v in enumerate(row, 1):
            _cell(ws, ws.max_row + 1 if j == 1 else ws.max_row, j, v)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = HEAD_FILL
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 27
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def export_full(campaign: dict, scans_dir: Path, names: list[str], ocr: dict, dest: Path) -> dict:
    """Écrit le dossier complet dans `dest` (fichier ZIP) — jamais en mémoire : il peut dépasser 500 Mo."""
    contracts = [evaluate(c) for c in campaign["contracts"]]
    saved = missing = 0
    with zipfile.ZipFile(dest, "w", zipfile.ZIP_STORED) as z:
        z.writestr("Global/Rque_contrat_Global.xlsx", make_workbook(contracts))
        for source, name in output_names(contracts).items():
            rows = [c for c in contracts if c["source"] == source]
            z.writestr(f"RAS/{name}", make_workbook([c for c in rows if not c["observations"]]))
            z.writestr(f"REAMRQUE/{name}", make_workbook([c for c in rows if c["observations"]]))
        for name, text in (campaign.get("inputs") or {}).items():
            z.writestr("Global/Inputs/" + re.sub(r"[^a-z0-9._-]", "_", name, flags=re.I), text)
        z.writestr("Global/campagne.json", json.dumps({"contracts": contracts, "inputs": campaign.get("inputs", {}),
                                                        "assetNames": names, "ocrRecords": ocr}, ensure_ascii=False, indent=2))
        if ocr:
            z.writestr("Global/Controle_pieces.xlsx", ocr_workbook(contracts, ocr))
            z.writestr("Global/OCR/resultats.json", json.dumps(ocr, ensure_ascii=False, indent=2))
        available_names = set(names)
        manifest = [["Numéro de contrat", "Qualité", "Abréviation", "Nom et prénom", "Pièce renommée", "Fichier original",
                     "Disponibilité"]]
        have: set[str] = set()
        for e in plan_pieces(contracts, names, campaign.get("aliases")):
            f = scans_dir / e["original"]
            ok = e["original"] in available_names and f.is_file()
            if ok:
                z.write(f, e["path"])
                saved += 1
                have.add(e["contract"])
            else:
                missing += 1
            manifest.append([e["contract"], e["role"], e["abbr"], e["person"], e["path"], e["original"],
                             "Fournie" if ok else "Référence introuvable"])
        for code in dict.fromkeys(c["number"] for c in contracts):
            if code not in have:
                z.writestr(f"PIECE SCANNEE/{_safe(code)}/AUCUNE_PIECE_FOURNIE.txt",
                           "Aucune pièce scannée disponible pour ce contrat. Les références introuvables sont conservées dans l’inventaire.")
        z.writestr("PIECE SCANNEE/Inventaire_pieces.csv",
                   "﻿" + "\r\n".join(";".join(_csv_cell(v) for v in r) for r in manifest))
        z.writestr("PIECE SCANNEE/LEGENDE.txt", "\n".join(f"{a} = {r}" for r, a in ROLE_CODES.items())
                   + "\nQUAL_A_CONFIRMER = rôle non établi.\nLes noms originaux sont conservés dans Inventaire_pieces.csv.")
    return {"savedPieces": saved, "missingPieces": missing}
