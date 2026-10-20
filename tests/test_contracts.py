"""Parité avec les tests de la plateforme v1 (control.mjs, downloads.mjs, scan-files.mjs, pieces.mjs)."""
import io
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.campaign_export import export_remarks, make_workbook, output_names, plan_pieces
from app.contracts import (COLUMNS, evaluate, export_row, from_snapshot, identity_remarks, parse_csv, unsigned)
from app.identity_check import inventory_aliases, resolve_scan

MINI = """### TABLE: contracts_new (2 lignes)
id;codeContract;type;village;duration;terrain_id;owner_id;beneficiary_id;numeroDemande
1;CT-1;PRET;Gbliglo;5;1;1;1;D1
2;CT-2;VENTE;Gbliglo;;2;2;2;
### TABLE: owners (2 lignes)
id;codeContract;name;idNumber;phoneNumber
1;CT-1;Jean Koffi;CI001;0102
2;CT-2;Awa Traore;CI002;
### TABLE: beneficiaries (2 lignes)
id;codeContract;name;idNumber;phoneNumber
1;CT-1;Yao Kouame;CI003;
2;CT-2;Paul Zadi;;
### TABLE: terrains (2 lignes)
id;codeContract;superficie
1;CT-1;2,5
2;CT-2;
### TABLE: contractSignatory (2 lignes)
id;codeContract;ownerName;ownerSignature;beneficiaryName;beneficiarySignature;ownerFingerprint
1;CT-1;Jean Koffi;oui;Yao Kouame;oui;["/a/b/scan1.jpg"]
2;CT-2;Awa Traore;oui;Paul Zadi;non;[]
"""


def _c():
    return {c["number"]: c for c in parse_csv(MINI, "src.csv")}


def test_parse_et_regles():
    rows = _c()
    c1, c2 = rows["CT-1"], rows["CT-2"]
    assert [a["role"] for a in c1["actors"]] == ["Propriétaire", "Bénéficiaire"]
    assert c1["actors"][0]["pieces"] == ["/a/b/scan1.jpg"] and c1["surface"] == "2,5"
    groups = {o["group"] for o in c1["observations"]}
    assert "Demande" not in groups and "Durée" not in groups and "Superficie" not in groups
    assert any("Pièce" in o["text"] or "pièce" in o["text"] for o in c1["observations"])   # contrôle en attente
    # VENTE : durée non exigée ; signature absente ; superficie manquante ; demande manquante
    assert not any(o["group"] == "Durée" for o in c2["observations"])
    texts = " | ".join(o["text"] for o in c2["observations"])
    assert "signature absente" in texts and "Superficie non rensegnée" in texts and "Numero de demande à preciser" in texts
    assert evaluate({**c2, "type": "PRET"})["observations"] and any(
        o["group"] == "Durée" for o in evaluate({**c2, "type": "PRET"})["observations"])
    assert not any(o["group"] == "Demande" for o in evaluate({**c2, "request": "000123"})["observations"])


def test_export_row_13_colonnes():
    r = export_row(_c()["CT-2"])
    assert len(COLUMNS) == len(r) == 13 and r[10] == "Numero de demande à preciser" and r[11] == "Superficie non rensegnée"
    assert unsigned(_c()["CT-1"]) == "RAS" and "Bénéficiaire" in unsigned(_c()["CT-2"])


def test_pourcentages_et_duree():
    base = {**_c()["CT-1"], "type": "PP_PLANTATION", "ownerShare": "60", "beneficiaryShare": "50", "request": "x"}
    assert any("Partage incohérent" in o["text"] for o in evaluate(base)["observations"])
    ok = evaluate({**base, "beneficiaryShare": "40"})
    assert not any(o["group"] == "Pourcentages" for o in ok["observations"])
    assert export_row({**base, "beneficiaryShare": "40"})[9] == "RAS"
    assert export_row(base)[9] == "Pourcentage superieur à 100%"
    assert export_row({**base, "duration": "0"})[8] == "Durée nulle"
    assert export_row({**base, "duration": "120"})[8] == "Durée invalide, sup à 99 ans"
    assert export_row({**base, "duration": "24", "raw": {"durationUnit": "mois"}})[8] == "24 mois"


def test_champs_obligatoires_inutilisables():
    c = {**_c()["CT-1"], "village": "???", "type": "VENTE", "request": "ras"}
    txt = identity_remarks(c) + "\n".join(export_row(c))
    assert "Village : valeur non exploitable" in txt and "Numéro de demande : valeur non exploitable" in txt


def test_excel_remarques(tmp_path):
    rows = [evaluate(c) for c in _c().values()]
    name, data, mime, n = export_remarks(rows)
    assert name == "Rque_contrat_src.xlsx" and n == 2 and data[:2] == b"PK"
    ws = load_workbook(io.BytesIO(data)).worksheets[0]
    assert ws.max_column == 13 and ws["K1"].value == "Num_demande" and ws["L1"].value == "Superficie"
    # deux sources -> ZIP avec Global ; noms nettoyés, collisions numérotées
    a = [{**rows[0], "source": "ALL_TABLES_EXPORT_EMILIENNE.csv"}, {**rows[1], "source": "EMILIENNE.csv"}]
    assert list(output_names(a).values()) == ["Rque_contrat_EMILIENNE.xlsx", "Rque_contrat_EMILIENNE_2.xlsx"]
    name, data, mime, _ = export_remarks(a)
    assert name == "Rque_contrat_Resultats.zip"
    assert "Global/Rque_contrat_Global.xlsx" in zipfile.ZipFile(io.BytesIO(data)).namelist()


def test_excel_pas_de_formule():
    c = {**_c()["CT-1"], "manual": "=1+1"}
    ws = load_workbook(io.BytesIO(make_workbook([evaluate(c)]))).worksheets[0]
    assert ws["M2"].data_type == "s" and "=1+1" in ws["M2"].value


def test_scans_et_inventaire():
    assert resolve_scan("C:\\scans\\piece.jpg", ["piece.jpg"]) == "piece.jpg"
    assert resolve_scan("piece.jpg", ["a/piece.jpg", "b/piece.jpg"]) is None
    assert resolve_scan("absente.jpg", ["autre.jpg"]) is None
    names = ["CT-001_PROP_JEAN_01.jpg"]
    al = inventory_aliases([{"Fichier original": "_Jean_123_0.jpg", "Pièce renommée": "PIECE SCANNEE/CT-001/CT-001_PROP_JEAN_01.jpg"}], names)
    assert resolve_scan("C:\\ancienne\\_Jean_123_0.jpg", names, al) == names[0]


def test_plan_pieces():
    c = {**_c()["CT-1"]}
    c["actors"] = [{**c["actors"][0], "pieces": ["scan.jpg"]}]
    p = plan_pieces([c], ["scan.jpg", "CT-375-009-002000-03-PRE_autre.jpg"])
    assert len(p) == 2 and p[0]["abbr"] == "PROP" and p[1]["abbr"] == "QUAL_A_CONFIRMER"
    assert p[0]["path"].startswith("PIECE SCANNEE/CT-1/") and p[0]["path"].endswith(".jpg")
    c["actors"] = [c["actors"][0], {**c["actors"][0], "role": "Bénéficiaire"}]
    assert plan_pieces([c], [])[0]["abbr"] == "QUAL_A_CONFIRMER"


def test_snapshot_historique():
    h = from_snapshot({"Tous les contrats": [{"Numéro de contrat": "C", "Fichier source": "s", "Type": "VENTE"}],
                       "Incohérences": [{"Numéro de contrat": "C", "Fichier source": "s", "Acteur": "PROP",
                                         "Champ contrôlé": "Nom signataire / acteur", "Statut": "Incohérence", "Détail": "A / B"}]})[0]
    assert "A / B" in "\n".join(export_row(h)) and any("A / B" in o["text"] for o in evaluate(h)["observations"])


REAL = Path("/tmp/claude-0/-home-claude/fbfdfc58-68c4-5b5d-9e87-7f851db22ba7/scratchpad/data/csv/ALL_TABLES_EXPORT_EMILIENNE.csv")


@pytest.mark.skipif(not REAL.exists(), reason="jeu de données réel absent")
def test_parite_donnees_reelles():
    rows = parse_csv(REAL.read_text(encoding="utf-8-sig"), "ALL_TABLES_EXPORT_EMILIENNE.csv")
    assert len(rows) == 190 and sum(len(c["actors"]) for c in rows) == 950
    assert sum(1 for c in rows for a in c["actors"] if not a["signed"]) == 0
    assert sum(1 for c in rows if any(o["group"] == "Pourcentages" for o in c["observations"])) == 10
    assert sum(1 for c in rows if any(o["group"] == "Identité" and "nom différent" in o["text"] for o in c["observations"])) == 1
