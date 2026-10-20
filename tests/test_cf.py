"""Contrôle pièce CF : jointures village / sous-préfecture, qualité, doublons, numéro, base des pièces validées."""
import io
import time
import zipfile

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import app.cf_api as cfapi
from app import cf, server
from app.config import Settings
from app.schemas import Extraction

CSV = """### TABLE: applicant (4 lignes)
id;applicantNumber;firstname;lastname;identityDocumentType;identityDocumentNumber;photoidentity
1;APT380-014-055001;GNASSE;Garcia;Carte nationale d'identité;CI005285407;["/x/380-014-055001_0.jpg","/x/380-014-055001_1.jpg"]
2;APT380-014-055002;KOUAKOU;Serge;Carte nationale d'identité;CI004878138;["/x/380-014-055002_0.jpg","/x/380-014-055002_1.jpg"]
3;APT380-013-000003;DIA;Claude;Carte nationale d'identité;CI005123675;["/x/380-013-000003_0.jpg","/x/380-013-000003_1.jpg"]
4;APT380-013-000004;KOKO;Felix;Carte nationale d'identité;CI004107757;["/x/380-013-000004_0.jpg","/x/380-013-000004_1.jpg"]
### TABLE: village (2 lignes)
id;codeVillage;label;codeSubPrefecture
1;380-014;Touagui 2;380
2;380-013;Touagui 1;380
### TABLE: subprefecture (1 lignes)
id;codeSubPrefecture;nameOfSubPrefect;label
1;380;MEAGUI;MEAGUI
### TABLE: request (1 lignes)
id;requestNumber;nomPrenomCE
1;380-014-055001;N'GUESSAN Regina
"""


def face(seed: int, w: int, h: int = 600, blur: bool = False):
    rng = np.random.default_rng(seed)
    img = np.full((h, w, 3), 235, np.uint8)
    for _ in range(60):
        p = tuple(int(v) for v in rng.integers(0, [w, h]))
        cv2.putText(img, "".join(rng.choice(list("ABCDEFGH012345"), 6)), p, cv2.FONT_HERSHEY_SIMPLEX, 0.8, (20, 20, 20), 2)
        cv2.rectangle(img, p, (p[0] + int(rng.integers(10, 80)), p[1] + int(rng.integers(10, 60))), (60, 90, 40), 2)
    if blur:
        img = cv2.GaussianBlur(img, (0, 0), 7)
    return img


def jpg(img) -> bytes:
    return cv2.imencode(".jpg", img)[1].tobytes()


class Fake:
    cache_id = "cf-fake"
    NUM = {800: "CI005285407", 810: "CI004878138", 820: "CI005123675", 830: "CI009999999"}

    def extract(self, img):
        w = img.size[0]
        recto = w in self.NUM
        return Extraction(lisible=True, type_piece="CNI", face="RECTO" if recto else "VERSO",
                          nom="X" if recto else None, numero_piece=self.NUM.get(w),
                          confiance={"nom": .95, "prenoms": .95, "numero_piece": .95})


def build(d):
    (d / "Images").mkdir(parents=True)
    (d / "ALL.csv").write_text(CSV, encoding="utf-8")
    W = lambda n, im: (d / "Images" / n).write_bytes(jpg(im))   # noqa: E731
    W("380-014-055001_0.jpg", face(1, 800)); W("380-014-055001_1.jpg", face(2, 801))                      # valide
    W("380-014-055002_0.jpg", face(3, 810)); W("380-014-055002_1.jpg", face(3, 810)[5:, 5:])             # recto x2
    W("380-013-000003_0.jpg", face(4, 820, blur=True)); W("380-013-000003_1.jpg", face(5, 821))          # flou
    W("380-013-000004_0.jpg", face(6, 830)); W("380-013-000004_1.jpg", face(7, 831))                     # n° différent
    W("scan_null_123_0.jpg", face(8, 840))                                                               # sans demande


def test_jointures():
    e = cf.parse_export(CSV)["applicants"]
    assert (e[0]["demande"], e[0]["code_village"], e[0]["village"], e[0]["sous_prefecture"]) == ("380-014-055001", "380-014", "Touagui 2", "MEAGUI")
    assert e[2]["village"] == "Touagui 1" and e[0]["ce"] == "N'GUESSAN Regina"


def test_prefixe_ci():
    assert cf.compare_cf_number("CI005285407", "1005285407", .5).status == "OK"
    assert cf.compare_cf_number("CI005285407", "CI005285401", .95).status == "PROCHE"
    assert cf.compare_cf_number("CI005285407", "CI009999999", .95).status == "ECART"


def test_session(tmp_path, monkeypatch):
    d = tmp_path / "s"
    build(d)
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    s = Settings(workers=1, cache_dir=tmp_path / "cache")
    data = cf.run_session(d, s, Fake())
    R = {r["demande"]: r for r in data["results"]}
    assert R["380-014-055001"]["statut"] == "VALIDE"
    assert R["380-014-055002"]["statut"] == "INVALIDE" and any("en double" in m for m in R["380-014-055002"]["motifs"])
    assert any("floue" in m for m in R["380-013-000003"]["motifs"]) and R["380-013-000003"]["statut"] == "INVALIDE"
    assert any("différent de la table applicant" in m for m in R["380-013-000004"]["motifs"])
    assert [o["fichier"] for o in data["orphans"]] == ["scan_null_123_0.jpg"]
    assert cf.store_valid(data["results"], d, "t") == 1
    rows = cf.valid_rows()
    assert len(rows) == 1 and rows[0]["village"] == "Touagui 2" and (cf.valid_root() / rows[0]["recto"]).is_file()
    wb = load_workbook(io.BytesIO(cf.invalid_workbook(data)))
    inv = wb["Invalidées"]
    assert inv.max_row == 4 and {inv.cell(r, 1).value for r in range(2, 5)} == {"380-014-055002", "380-013-000003", "380-013-000004"}
    assert "Motif(s)" in [c.value for c in inv[1]]


def test_api(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(cfapi, "get_extractor", lambda s, b=None: Fake())
    monkeypatch.setattr(cfapi, "get_settings", lambda: Settings(workers=1, cache_dir=tmp_path / "cache"))
    src = tmp_path / "src"
    build(src)
    c = TestClient(server.app)
    sid = c.post("/api/cf", data={"title": "t"}).json()["id"]
    files = [("files", ("ALL.csv", (src / "ALL.csv").read_bytes()))] + [("files", (p.name, p.read_bytes())) for p in (src / "Images").iterdir()]
    assert c.post(f"/api/cf/{sid}/files", files=files).json()["total_images"] == 9
    assert c.post(f"/api/cf/{sid}/run", data={"backend": "tesseract"}).status_code == 200
    for _ in range(100):
        st = c.get(f"/api/cf/{sid}/status").json()
        if st["status"] != "running":
            break
        time.sleep(0.1)
    assert st["status"] == "done", st
    r = c.get(f"/api/cf/{sid}/results").json()
    assert r["summary"]["VALIDE"] == 1 and r["summary"]["INVALIDE"] == 3 and r["summary"]["STOCKEES"] == 1
    assert c.get(f"/api/cf/{sid}/export/invalides.xlsx").content[:2] == b"PK"
    assert len(c.get("/api/cf/valid/list").json()) == 1
    z = zipfile.ZipFile(io.BytesIO(c.get("/api/cf/valid/export.zip").content))
    assert "Base_pieces_validees.xlsx" in z.namelist() and any(n.startswith("pieces/") for n in z.namelist())
    assert c.get(f"/api/cf/{sid}/scan/..%2Fx").status_code in (404, 400)
    assert "Contrôle pièce CF" in c.get("/cf").text
